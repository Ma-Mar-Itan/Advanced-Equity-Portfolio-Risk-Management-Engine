"""Turn slide_notes.md into a plain PowerPoint deck (QFI_V2_slides.pptx).

No styling on purpose: default template, title + bullet points, chart pictures,
speaker notes. Run from the V2 folder:  python -m tools.build_deck
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
NOTES_FILE = ROOT / "slide_notes.md"
DECK_FILE = ROOT / "QFI_V2_slides.pptx"
DECK_TITLE = "Equity Portfolio Risk"
DECK_SUBTITLE = "QFI Project #3: how much could we lose, and can we trust the answer?"


@dataclass
class Slide:
    """One slide as written in slide_notes.md."""

    title: str
    bullets: list[str] = field(default_factory=list)
    image: str | None = None
    say: str = ""


def parse_notes(text: str) -> list[Slide]:
    """Read slide_notes.md: '### ' starts a slide, '- ' adds a bullet, 'Image:' and 'Say:' set the rest."""
    slides: list[Slide] = []
    for line in text.splitlines():
        line = line.strip()
        if line.startswith("### "):
            slides.append(Slide(line[4:].strip()))
        elif not slides:
            continue
        elif line.startswith("- "):
            slides[-1].bullets.append(line[2:].strip())
        elif line.startswith("Image:"):
            slides[-1].image = line[6:].strip()
        elif line.startswith("Say:"):
            slides[-1].say = line[4:].strip()
    return slides


def set_title(slide, text: str) -> None:
    """Write the slide title at one size on every slide (32pt, so two lines still fit)."""
    slide.shapes.title.text = text
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)


def add_text_slide(prs: Presentation, s: Slide) -> None:
    """A 'Title and Content' slide: title plus bullet points."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, s.title)
    body = slide.placeholders[1].text_frame
    for i, bullet in enumerate(s.bullets):
        para = body.paragraphs[0] if i == 0 else body.add_paragraph()
        para.text = bullet
        para.font.size = Pt(20)
    slide.notes_slide.notes_text_frame.text = s.say


def add_image_slide(prs: Presentation, s: Slide) -> None:
    """A 'Title Only' slide: title, the chart, and the bullet points underneath."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(slide, s.title)
    slide.shapes.add_picture(str(ROOT / s.image), Inches(0.7), Inches(1.6), width=Inches(8.6))
    box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.95)).text_frame
    box.word_wrap = True
    for i, bullet in enumerate(s.bullets):
        para = box.paragraphs[0] if i == 0 else box.add_paragraph()
        para.text = "• " + bullet
        para.font.size = Pt(14)
    slide.notes_slide.notes_text_frame.text = s.say


def build_deck(slides: list[Slide]) -> Presentation:
    """A title slide followed by one slide per entry in the notes."""
    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = DECK_TITLE
    cover.placeholders[1].text = DECK_SUBTITLE
    for s in slides:
        (add_image_slide if s.image else add_text_slide)(prs, s)
    return prs


def main() -> None:
    """Rebuild the deck from the notes file."""
    slides = parse_notes(NOTES_FILE.read_text(encoding="utf-8"))
    for s in slides:
        if s.image and len(s.bullets) > 2:
            print(f"Warning: '{s.title}' has {len(s.bullets)} bullets; image slides fit 2.")
    build_deck(slides).save(DECK_FILE)
    print(f"{len(slides) + 1} slides written to {DECK_FILE}")


if __name__ == "__main__":
    main()
