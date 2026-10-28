"""Build a short Word glossary for the QFI project learning outcomes."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "QFI_Expected_Learning_Outcomes_Glossary.docx"


CONCEPTS = [
    (
        "Equity portfolio risk modeling",
        "The process of measuring how a group of stocks may gain or lose value. It combines each stock's return, portfolio weight, volatility, and relationship with the other holdings to estimate total portfolio risk.",
    ),
    (
        "Historical Simulation VaR",
        "A Value at Risk method that ranks actual past portfolio returns from worst to best and selects a loss cutoff at a chosen confidence level. It is transparent, but it assumes the selected history is informative about future risk.",
    ),
    (
        "Volatility scaling",
        "A method that resizes past returns to reflect today's market volatility. Historical shocks are increased when current risk is higher and reduced when current conditions are calmer.",
    ),
    (
        "EWMA volatility",
        "Exponentially Weighted Moving Average volatility gives more importance to recent market movements and gradually less importance to older ones. The lambda parameter controls how quickly old information loses influence.",
    ),
    (
        "GARCH modeling",
        "A statistical approach that models changing volatility through time. It captures volatility clustering, where large market movements tend to be followed by more large movements and calm periods tend to persist.",
    ),
    (
        "GJR-GARCH leverage effect",
        "An extension of GARCH that allows negative returns to raise future volatility more than equally large positive returns. This asymmetric response is useful for equities because bad news often creates a stronger increase in uncertainty.",
    ),
    (
        "Expected Shortfall",
        "The average loss on days when the loss is already worse than the VaR threshold. VaR shows where the severe-loss region begins; Expected Shortfall describes the average severity inside that region.",
    ),
    (
        "Backtesting exceptions",
        "An exception occurs when the actual portfolio loss is greater than the VaR forecast made before that day. Backtesting checks whether the number and timing of these exceptions are consistent with the selected confidence level.",
    ),
    (
        "Kupiec test",
        "A statistical test of exception frequency. It checks whether the total number of VaR exceptions is plausible for the chosen confidence level, but it does not test whether exceptions are clustered.",
    ),
    (
        "Christoffersen test",
        "A backtesting test that examines whether VaR exceptions occur independently rather than in clusters. Conditional coverage combines this timing test with the exception-frequency test.",
    ),
    (
        "Basel Traffic Light framework",
        "A regulatory-style classification that places a VaR model in a green, yellow, or red zone according to its exception count. Green is favorable, but the traffic light alone does not prove that the model is fully reliable.",
    ),
    (
        "Stress testing",
        "A what-if analysis that estimates portfolio losses under severe historical or hypothetical scenarios. Stress results are conditional estimates, not forecasts or statements about how likely an event is.",
    ),
    (
        "Risk dashboard development",
        "The design of an interactive application that brings portfolio inputs, risk estimates, charts, backtesting results, and stress scenarios together. A useful dashboard makes the analysis understandable, comparable, and easy to monitor.",
    ),
]


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def remove_paragraph_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title._element.get_or_add_pPr()
    title_p_bdr = title_p_pr.find(qn("w:pBdr"))
    if title_p_bdr is not None:
        title_p_pr.remove(title_p_bdr)

    p = doc.add_paragraph(style="Title")
    p.add_run("Expected Learning Outcomes")
    p.paragraph_format.space_after = Pt(8)
    remove_paragraph_border(p)

    intro = doc.add_paragraph(
        "This glossary summarizes the main concepts developed in the equity portfolio risk project. "
        "Each definition focuses on the practical idea and how it supports risk measurement or decision-making."
    )
    intro.paragraph_format.space_after = Pt(12)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.4)
    table.columns[1].width = Inches(4.65)

    header = table.rows[0]
    repeat_header(header)
    keep_row_together(header)
    for idx, text in enumerate(("Concept", "Plain-language explanation")):
        cell = header.cells[idx]
        shade_cell(cell, "1F4E79")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(text)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, (concept, explanation) in enumerate(CONCEPTS, start=1):
        row = table.add_row()
        keep_row_together(row)
        cells = row.cells
        cells[0].width = Inches(2.4)
        cells[1].width = Inches(4.65)
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        if row_index % 2 == 0:
            shade_cell(cells[0], "EEF3F7")
            shade_cell(cells[1], "EEF3F7")

        concept_p = cells[0].paragraphs[0]
        concept_r = concept_p.add_run(concept)
        concept_r.bold = True
        concept_r.font.name = "Aptos"
        concept_r.font.size = Pt(10.2)
        concept_r.font.color.rgb = RGBColor(25, 61, 92)

        explanation_p = cells[1].paragraphs[0]
        explanation_r = explanation_p.add_run(explanation)
        explanation_r.font.name = "Aptos"
        explanation_r.font.size = Pt(10.2)

    doc.core_properties.title = "Expected Learning Outcomes"
    doc.core_properties.subject = "Plain-language glossary for the QFI equity portfolio risk project"
    doc.core_properties.author = "Marwan Itani"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
